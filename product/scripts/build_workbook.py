#!/usr/bin/env python3
"""Rebuild GL_Hack_Adoption_Quant_Model.xlsx from sheets/*.csv."""
from __future__ import annotations

import csv
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill, Border, Side
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parents[1]
SHEETS = ROOT / "sheets"
OUT = ROOT / "GL_Hack_Adoption_Quant_Model.xlsx"

VOID = "0B0F0C"
PHOSPHOR = "3DFF8A"
DIM = "163226"
PALE = "E8F2EC"
WHITE = "F4FBF6"
thin = Border(
    left=Side(style="thin", color="1F3A2E"),
    right=Side(style="thin", color="1F3A2E"),
    top=Side(style="thin", color="1F3A2E"),
    bottom=Side(style="thin", color="1F3A2E"),
)
ORDER = [
    "00_Cover",
    "01_Method",
    "02_Assumptions",
    "03_Product",
    "04_Segments",
    "05_Analogs",
    "06_Results",
    "07_Waterfall",
    "08_Sensitivity",
]


def load_csv(name: str) -> list[list[str]]:
    path = SHEETS / f"{name}.csv"
    with path.open(newline="", encoding="utf-8") as f:
        return [row for row in csv.reader(f)]


def write_sheet(wb: Workbook, name: str, first: bool) -> None:
    rows = load_csv(name)
    ws = wb.active if first else wb.create_sheet(name)
    if first:
        ws.title = name
    max_cols = max((len(r) for r in rows), default=1)
    for r_i, row in enumerate(rows, 1):
        for c_i, val in enumerate(row, 1):
            cell = ws.cell(r_i, c_i, val)
            cell.font = Font(name="Space Grotesk", size=11)
            cell.alignment = Alignment(wrap_text=True, vertical="center")
            cell.border = thin
            if r_i == 1:
                cell.font = Font(name="Space Grotesk", size=14, bold=True, color=WHITE)
                cell.fill = PatternFill("solid", fgColor=VOID)
            elif r_i == 2:
                cell.fill = PatternFill("solid", fgColor=DIM)
                cell.font = Font(name="Space Mono", size=11, color=PHOSPHOR)
            elif r_i % 2 == 0:
                cell.fill = PatternFill("solid", fgColor=PALE)
        ws.row_dimensions[r_i].height = 18 if r_i > 2 else 22
    for col in range(1, max_cols + 1):
        ws.column_dimensions[get_column_letter(col)].width = 18 if col > 1 else 28
    ws.freeze_panes = "A4"
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToPage = True
    ws.page_setup.fitToWidth = 1
    ws.sheet_view.showGridLines = False
    ws.oddFooter.right.text = name


def main() -> None:
    wb = Workbook()
    for i, name in enumerate(ORDER):
        write_sheet(wb, name, first=(i == 0))
    wb.save(OUT)
    print("wrote", OUT)


if __name__ == "__main__":
    main()
